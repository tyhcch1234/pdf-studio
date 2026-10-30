"""Browser smoke test for the three tools (requires playwright + chromium).

    python3 tests/ui_test.py [base_url]
"""
from __future__ import annotations

import os
import sys

import pymupdf
from playwright.sync_api import sync_playwright

BASE = sys.argv[1] if len(sys.argv) > 1 else "http://127.0.0.1:8000"
HERE = os.path.dirname(os.path.abspath(__file__))
SAMPLES = os.path.join(HERE, "samples")
OUT = os.path.join(HERE, "out")
os.makedirs(OUT, exist_ok=True)

fails: list[str] = []


def check(cond: bool, msg: str) -> None:
    print(("  PASS  " if cond else "  FAIL  ") + msg)
    if not cond:
        fails.append(msg)


def main() -> int:
    with sync_playwright() as pw:
        browser = pw.chromium.launch()
        page = browser.new_page(viewport={"width": 1280, "height": 1000})
        errors: list[str] = []
        page.on("pageerror", lambda e: errors.append(str(e)))
        page.on("console", lambda m: errors.append(m.text) if m.type == "error" else None)

        page.goto(BASE, wait_until="networkidle")
        page.wait_for_timeout(1200)
        check("PDF Studio" in page.title(), "page loads with title")
        caps = page.inner_text("#cap-ocr")
        check("就緒" in caps, f"OCR engine reported ready ({caps})")

        # ---------- 1. PDF -> Word ----------
        page.set_input_files(
            "#input-convert",
            [os.path.join(SAMPLES, "digital.pdf"), os.path.join(SAMPLES, "scanned.pdf")],
        )
        page.wait_for_timeout(400)
        check(page.locator("#list-convert .frow").count() == 2, "two files listed for conversion")
        page.click("#btn-convert")
        page.wait_for_selector("#results-convert .result", timeout=180000)
        page.wait_for_timeout(1500)
        page.screenshot(path=os.path.join(OUT, "ui_convert.png"), full_page=True)

        cards = page.locator("#results-convert .result")
        check(cards.count() >= 2, f"conversion results rendered ({cards.count()} cards)")
        docx_card = page.locator("#results-convert .result", has_text=".docx").first
        docx_links = docx_card.locator(".btn")
        with page.expect_download(timeout=60000) as dl:
            docx_links.first.click()
        path = os.path.join(OUT, "ui_download.docx")
        dl.value.save_as(path)
        import docx

        d = docx.Document(path)
        text = "\n".join(p.text for p in d.paragraphs)
        check("Measurement Report" in text, "docx keeps the digital text layer")
        check(len(d.inline_shapes) >= 1, "docx keeps the picture")
        ns = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
        has_math = any(p._p.findall(f".//{ns}oMath") for p in d.paragraphs)
        print(f"  note   equations in downloaded file: {has_math}")

        # ---------- 1b. equations must survive as native OMML ----------
        page.click("#btn-clear-convert")
        page.wait_for_timeout(300)
        page.set_input_files("#input-convert", [os.path.join(SAMPLES, "math_sample.pdf")])
        page.wait_for_timeout(300)
        page.click("#btn-convert")
        page.wait_for_selector("#results-convert .result", timeout=180000)
        page.wait_for_timeout(1200)
        card = page.locator("#results-convert .result", has_text=".docx").first
        with page.expect_download(timeout=60000) as dlm:
            card.locator(".btn").first.click()
        math_path = os.path.join(OUT, "ui_math.docx")
        dlm.value.save_as(math_path)
        md = docx.Document(math_path)
        n_math = sum(1 for p in md.paragraphs if p._p.findall(f".//{ns}oMath"))
        sup = sum(1 for p in md.paragraphs for e in p._p.iter(f"{ns}sSup"))
        check(n_math == 2, f"both display equations became Word equations (got {n_math})")
        check(sup >= 2, f"equation structure keeps superscripts (got {sup})")

        # ---------- 2. merge ----------
        page.click('.tab[data-panel="merge"]')
        page.set_input_files(
            "#input-merge",
            [os.path.join(SAMPLES, "a3_landscape.pdf"), os.path.join(SAMPLES, "digital.pdf")],
        )
        page.wait_for_timeout(1200)
        check(page.locator("#list-merge .frow").count() == 2, "two files listed for merge")
        page.click("#btn-merge")
        page.wait_for_selector("#results-merge .result", timeout=60000)
        with page.expect_download(timeout=30000) as dl2:
            page.locator("#results-merge .result .btn").first.click()
        merge_path = os.path.join(OUT, "ui_merged.pdf")
        dl2.value.save_as(merge_path)
        merged = pymupdf.open(merge_path)
        check(merged.page_count == 2, f"merged pdf has 2 pages (got {merged.page_count})")
        check(
            abs(merged[0].rect.width - 1191) < 2 and abs(merged[0].rect.height - 842) < 2,
            "merged pdf keeps the A3 first page",
        )
        page.screenshot(path=os.path.join(OUT, "ui_merge.png"), full_page=True)

        # ---------- 3. A3 -> A4 ----------
        page.click('.tab[data-panel="split"]')
        page.set_input_files("#input-split", [os.path.join(SAMPLES, "a3_landscape.pdf")])
        page.wait_for_timeout(2000)
        check(page.locator("#previews-split .pv").count() == 1, "preview rendered for A3 page")
        check(page.locator("#previews-split .cut.v").count() == 1, "cut line drawn vertically")
        page.screenshot(path=os.path.join(OUT, "ui_split.png"), full_page=True)
        page.click("#btn-split")
        page.wait_for_selector("#results-split .result", timeout=60000)
        with page.expect_download(timeout=30000) as dl3:
            page.locator("#results-split .result .btn").first.click()
        split_path = os.path.join(OUT, "ui_split.pdf")
        dl3.value.save_as(split_path)
        sp = pymupdf.open(split_path)
        check(sp.page_count == 2, f"split pdf has 2 pages (got {sp.page_count})")
        w0, h0 = sp[0].rect.width, sp[0].rect.height
        check(abs(w0 - 595) < 3 and abs(h0 - 842) < 3, f"A4 portrait output ({w0:.0f}x{h0:.0f})")
        check("LEFT" in sp[0].get_text(), "first page holds the left half")
        check("RIGHT" in sp[1].get_text(), "second page holds the right half")

        if errors:
            print("  console errors:")
            for e in errors[:8]:
                print("   -", e[:160])
        check(not errors, "no javascript errors")
        browser.close()

    print("\nFAILED:" if fails else "\nALL PASSED")
    for f in fails:
        print("  -", f)
    return 1 if fails else 0


if __name__ == "__main__":
    raise SystemExit(main())

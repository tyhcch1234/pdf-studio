"""Layout model -> .docx writer (text, pictures and native equations)."""
from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.shared import Pt, RGBColor

from .extract import Element, PageData
from .omml import latex_to_omml, omml_from_latex_fallback

MARGIN = 14.0  # points

# XML 1.0 forbids most control characters; PDF fonts can yield them.
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")


def _xml_safe(text: str) -> str:
    text = _ILLEGAL.sub("", text)
    return text.encode("utf-8", "ignore").decode("utf-8")


def _setup_section(section, page: PageData) -> None:
    section.page_width = Pt(page.width)
    section.page_height = Pt(page.height)
    section.left_margin = Pt(MARGIN)
    section.right_margin = Pt(MARGIN)
    section.top_margin = Pt(MARGIN)
    section.bottom_margin = Pt(MARGIN)
    section.header_distance = Pt(6)
    section.footer_distance = Pt(6)


def _new_paragraph(doc, gap_before: float, indent: float):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(max(0.0, min(gap_before, 26.0)))
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    pf.left_indent = Pt(max(0.0, indent))
    pf.widow_control = False
    return p


def _add_runs(p, el: Element) -> None:
    for run in el.runs:
        if not run.text:
            continue
        safe = _xml_safe(run.text)
        if not safe:
            continue
        r = p.add_run(safe)
        r.font.size = Pt(max(5.0, round(run.size, 1)))
        r.bold = run.bold
        r.italic = run.italic
        if run.superscript:
            r.font.superscript = True
        if run.subscript:
            r.font.subscript = True
        if run.color != (0, 0, 0):
            r.font.color.rgb = RGBColor(*run.color)


def _add_equation(p, latex: str, display: bool = True) -> bool:
    xml = latex_to_omml(latex, display=display) or omml_from_latex_fallback(latex)
    if not xml:
        return False
    try:
        p._p.append(parse_xml(xml))
        return True
    except Exception:
        return False


def _add_picture(p, png: bytes, width_pt: float) -> bool:
    try:
        p.add_run().add_picture(io.BytesIO(png), width=Pt(max(4.0, width_pt)))
        return True
    except Exception:
        return False


def write_docx(pages: list[PageData], out_path: str, title: str = "") -> dict:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(0)

    stats = {
        "pages": len(pages),
        "ocr_pages": 0,
        "text_lines": 0,
        "images": 0,
        "equations": 0,
        "equations_as_image": 0,
        "warnings": [],
    }

    _setup_section(doc.sections[0], pages[0])
    prev_w = pages[0].width
    prev_h = pages[0].height

    for pi, page in enumerate(pages):
        if pi > 0:
            if abs(page.width - prev_w) > 1 or abs(page.height - prev_h) > 1:
                sec = doc.add_section(WD_SECTION.NEW_PAGE)
                _setup_section(sec, page)
                prev_w, prev_h = page.width, page.height
            else:
                doc.add_page_break()
        if page.ocr_used:
            stats["ocr_pages"] += 1

        prev_bottom = MARGIN
        for el in page.elements:
            gap = el.bbox[1] - prev_bottom
            indent = el.bbox[0] - MARGIN
            width_pt = el.bbox[2] - el.bbox[0]

            if el.kind == "text":
                if not any(run.text.strip() for run in el.runs):
                    continue
                p = _new_paragraph(doc, gap, indent)
                _add_runs(p, el)
                stats["text_lines"] += 1

            elif el.kind == "formula":
                p = _new_paragraph(doc, gap, indent)
                if el.latex and _add_equation(p, el.latex, display=True):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.left_indent = Pt(0)
                    stats["equations"] += 1
                elif el.png and _add_picture(p, el.png, width_pt):
                    stats["equations_as_image"] += 1
                    stats["warnings"].append(
                        f"page {page.index + 1}: formula kept as picture (no LaTeX)"
                    )
                else:
                    stats["warnings"].append(
                        f"page {page.index + 1}: formula region could not be recovered"
                    )

            elif el.kind == "image" and el.png:
                p = _new_paragraph(doc, gap, indent)
                if _add_picture(p, el.png, width_pt):
                    stats["images"] += 1

            prev_bottom = el.bbox[3]

    if title:
        doc.core_properties.title = title
    doc.save(out_path)

    # de-duplicate warnings, keep the report readable
    seen = set()
    unique = []
    for w in stats["warnings"]:
        key = w.split(":", 1)[-1]
        if key in seen:
            continue
        seen.add(key)
        unique.append(w)
    stats["warnings"] = unique[:20]
    return stats
